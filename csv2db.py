"""Explicit-schema spreadsheet importer and restricted read-only SQLite queries."""
import argparse
import csv
import hashlib
import json
import math
import os
import sqlite3
import tempfile
import time
from contextlib import closing
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path


def quote(name):
    if not isinstance(name, str) or not name or '\x00' in name:
        raise ValueError('标识符必须是非空字符串且不能含 NUL')
    return '"' + name.replace('"', '""') + '"'


def file_hash(path):
    digest = hashlib.sha256()
    with path.open('rb') as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b''):
            digest.update(block)
    return digest.hexdigest()


def convert(value, column):
    if value is None or value == '':
        if column.get('required', False):
            raise ValueError('必填字段为空')
        return None
    kind = column['type']
    if kind == 'text':
        if not isinstance(value, str):
            raise ValueError('文本字段包含非文本值，请明确转换规则，避免编号精度或前导零丢失')
        return value
    if kind == 'date':
        if isinstance(value, datetime):
            if value.time().isoformat() != '00:00:00':
                raise ValueError('日期字段含时间，不能静默丢弃')
            return value.date().isoformat()
        if isinstance(value, date):
            return value.isoformat()
        return date.fromisoformat(str(value)).isoformat()
    if isinstance(value, bool):
        raise ValueError('布尔值不能用作数值')
    try:
        number = Decimal(str(value))
        if not number.is_finite():
            raise ValueError('数值必须有限')
        if kind == 'money':
            number *= 100
        if kind in ('integer', 'money'):
            if number != number.to_integral_value():
                raise ValueError('整数含小数或金额超过两位小数；需要明确舍入规则')
            result = int(number)
            if not -(2**63) <= result < 2**63:
                raise ValueError('数值超过 SQLite INTEGER 范围')
            return result
        result = float(number)
        if not math.isfinite(result):
            raise ValueError('数值超出 REAL 范围')
        return result
    except InvalidOperation as exc:
        raise ValueError('非法数值') from exc


def source_rows(path, config):
    """Yield (source row, values); fail on formulas rather than trust stale caches."""
    if path.suffix.lower() == '.csv':
        with path.open(encoding=config.get('encoding', 'utf-8-sig'), newline='') as handle:
            for number, row in enumerate(csv.reader(handle), 1):
                yield number, row
    elif path.suffix.lower() in ('.xlsx', '.xlsm'):
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            sheet_name = config.get('sheet')
            if sheet_name is None:
                if len(workbook.sheetnames) != 1:
                    raise ValueError('多工作表文件必须指定 sheet')
                sheet_name = workbook.sheetnames[0]
            sheet = workbook[sheet_name]
            for number, row in enumerate(sheet.iter_rows(), 1):
                if any(cell.data_type in ('f', 'e') for cell in row):
                    raise ValueError(f'源行 {number} 含公式或 Excel 错误；请先提供已核验的数值副本')
                yield number, [cell.value for cell in row]
        finally:
            workbook.close()
    else:
        raise ValueError('仅支持 .csv、.xlsx、.xlsm')


def check_replace_target(source, output):
    if source == output:
        raise ValueError('输出不能与源文件相同')
    for suffix in ('-wal', '-shm', '-journal'):
        if Path(str(output) + suffix).exists():
            raise ValueError('目标数据库存在事务配套文件，请关闭使用它的程序后再替换')
    if output.exists() and not output.is_file():
        raise ValueError('输出必须是普通文件')


def saved_config(database):
    """Read only the importer-owned metadata; never infer types from SQL tables."""
    try:
        with closing(sqlite3.connect(database.as_uri() + '?mode=ro', uri=True)) as conn:
            rows = conn.execute('SELECT report_json FROM _import_metadata LIMIT 2').fetchall()
        if len(rows) != 1:
            raise ValueError('导入元数据必须只有一条记录')
        report = json.loads(rows[0][0])
        config = report['config']
        if (not isinstance(config, dict) or not isinstance(config.get('table'), str)
                or not isinstance(config.get('columns'), list) or not config['columns']):
            raise ValueError('保存的字段配置格式无效')
        for column in config['columns']:
            if not isinstance(column, dict) or any(
                    not isinstance(column.get(key), str) for key in ('source', 'name', 'type')):
                raise ValueError('保存的字段配置格式无效')
        return config
    except (sqlite3.Error, ValueError, KeyError, TypeError) as exc:
        raise ValueError('无法读取数据库中保存的字段配置，请显式提供 --schema') from exc


def import_file(source, output, config=None, replace=False):
    source, output = Path(source).resolve(), Path(output).resolve()
    if source == output:
        raise ValueError('输出不能与源文件相同')
    if replace:
        check_replace_target(source, output)
    elif output.exists():
        raise FileExistsError(f'输出已存在：{output}')
    reused_config = config is None
    if reused_config:
        if not replace or not output.is_file():
            raise ValueError('首次导入必须提供 --schema；已有数据库使用 --replace 时可省略')
        config = saved_config(output)
    table = config['table']
    if table.lower().startswith(('sqlite_', '_import_')):
        raise ValueError('表名使用了保留前缀')
    quote(table)
    columns = config['columns']
    if not columns:
        raise ValueError('至少需要一个字段')
    names = [c['name'] for c in columns]
    if len({n.casefold() for n in names}) != len(names) or '_source_row' in {n.casefold() for n in names}:
        raise ValueError('字段名重复或使用了保留字段')
    types = {'text': 'TEXT', 'date': 'TEXT', 'integer': 'INTEGER', 'money': 'INTEGER', 'real': 'REAL'}
    for col in columns:
        quote(col['name'])
        if col['type'] not in types:
            raise ValueError('未知字段类型')
    header_row = config.get('header_row', 1)
    if not isinstance(header_row, int) or header_row < 1:
        raise ValueError('header_row 必须是正整数')
    output.parent.mkdir(parents=True, exist_ok=True)
    fd, temporary = tempfile.mkstemp(prefix='.import-', suffix='.db', dir=output.parent)
    os.close(fd)
    report = {'source': str(source), 'table': table, 'mapping': columns,
              'read': 0, 'imported': 0, 'excluded_blank': 0,
              'formula_policy': 'reject', 'hidden_rows': 'included',
              'imported_at': datetime.now().astimezone().isoformat(),
              'sqlite_version': sqlite3.sqlite_version, 'config': config,
              'publish_mode': 'replace' if replace else 'create',
              'config_source': 'database' if reused_config else 'explicit'}
    try:
        before_hash = file_hash(source)
        expected = {c['name']: [0, 0] for c in columns if c['type'] in ('money', 'integer')}
        with closing(sqlite3.connect(temporary)) as conn:
            with conn:
                definitions = [quote(c['name']) + ' ' + types[c['type']] +
                               (' NOT NULL' if c.get('required') else '') for c in columns]
                definitions.append('"_source_row" INTEGER NOT NULL')
                conn.execute(f'CREATE TABLE {quote(table)} ({", ".join(definitions)})')
                primary = config.get('unique_key', [])
                if primary:
                    if not set(primary) <= set(names):
                        raise ValueError('unique_key 包含未知字段')
                    conn.execute(f'CREATE UNIQUE INDEX {quote(table + "_unique")} ON '
                                 f'{quote(table)} ({", ".join(quote(n) for n in primary)})')
                header = None
                batch = []
                insert = f'INSERT INTO {quote(table)} VALUES ({",".join("?" for _ in definitions)})'
                for number, values in source_rows(source, config):
                    if number < header_row:
                        continue
                    if number == header_row:
                        header = values
                        for c in columns:
                            if header.count(c['source']) != 1:
                                raise ValueError(f'源字段缺失或重名：{c["source"]}')
                        continue
                    report['read'] += 1
                    if all(v is None or v == '' for v in values):
                        report['excluded_blank'] += 1
                        continue
                    record = []
                    for col in columns:
                        index = header.index(col['source'])
                        value = values[index] if index < len(values) else None
                        try:
                            record.append(convert(value, col))
                        except (ValueError, TypeError) as exc:
                            raise ValueError(f'源行 {number}，字段 {col["source"]}：{exc}') from exc
                    if primary and any(record[names.index(n)] is None for n in primary):
                        raise ValueError(f'源行 {number}：唯一业务键为空')
                    batch.append(record + [number])
                    for index, col in enumerate(columns):
                        if col['name'] in expected and record[index] is not None:
                            expected[col['name']][0] += record[index]
                            expected[col['name']][1] += 1
                    report['imported'] += 1
                    if len(batch) >= 1000:
                        conn.executemany(insert, batch)
                        batch.clear()
                if header is None:
                    raise ValueError('找不到表头行')
                conn.executemany(insert, batch)
                if report['imported'] == 0:
                    raise ValueError('没有可导入的数据')
                if file_hash(source) != before_hash:
                    raise ValueError('导入过程中源文件发生变化')
                report['source_sha256'] = before_hash
                count = conn.execute(f'SELECT COUNT(*) FROM {quote(table)}').fetchone()[0]
                assert count == report['imported']
                assert report['read'] == count + report['excluded_blank']
                report['integrity_check'] = conn.execute('PRAGMA integrity_check').fetchone()[0]
                if report['integrity_check'] != 'ok':
                    raise ValueError('数据库完整性检查失败')
                report['numeric_checks'] = {}
                for c in columns:
                    if c['type'] in ('money', 'integer', 'real'):
                        report['numeric_checks'][c['name']] = conn.execute(
                            f'SELECT SUM({quote(c["name"])}), COUNT({quote(c["name"])}) '
                            f'FROM {quote(table)}').fetchone()
                        if c['name'] in expected:
                            total, populated = expected[c['name']]
                            if (total if populated else None, populated) != report['numeric_checks'][c['name']]:
                                raise ValueError(f'字段 {c["name"]} 汇总对账失败')
                report['source_numeric_checks'] = expected
                conn.execute('CREATE TABLE _import_metadata (report_json TEXT NOT NULL)')
                conn.execute('INSERT INTO _import_metadata VALUES (?)',
                             (json.dumps(report, ensure_ascii=False),))
            # Transaction committed before publishing.
        if replace:
            check_replace_target(source, output)
            os.replace(temporary, output)  # Same directory: atomic snapshot replacement.
        else:
            os.link(temporary, output)  # Atomic, refuses overwriting an existing destination.
        return report
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)


def query_file(database, sql, parameters=None, limit=1000, timeout=5):
    if limit < 1 or timeout <= 0:
        raise ValueError('limit 与 timeout 必须大于零')
    uri = Path(database).resolve().as_uri() + '?mode=ro'
    with closing(sqlite3.connect(uri, uri=True)) as conn:
        conn.execute('PRAGMA query_only=ON')
        allowed = {sqlite3.SQLITE_SELECT, sqlite3.SQLITE_READ,
                   sqlite3.SQLITE_FUNCTION, sqlite3.SQLITE_RECURSIVE}
        def authorize(action, arg1, arg2, db, trigger):
            if action == sqlite3.SQLITE_FUNCTION and str(arg2).lower() == 'load_extension':
                return sqlite3.SQLITE_DENY
            return sqlite3.SQLITE_OK if action in allowed else sqlite3.SQLITE_DENY
        conn.set_authorizer(authorize)
        deadline = time.monotonic() + timeout
        conn.set_progress_handler(lambda: int(time.monotonic() > deadline), 1000)
        try:
            cursor = conn.execute(sql, parameters if parameters is not None else {})
        except sqlite3.Warning as exc:
            raise ValueError('每次仅允许执行一条查询') from exc
        rows = cursor.fetchmany(limit + 1)
        return {'columns': [c[0] for c in cursor.description], 'rows': rows[:limit],
                'truncated': len(rows) > limit, 'sql': sql,
                'parameters': parameters if parameters is not None else {}}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    commands = parser.add_subparsers(dest='command', required=True)
    imp = commands.add_parser('import')
    imp.add_argument('source'); imp.add_argument('output')
    imp.add_argument('--schema', help='字段配置 JSON；更新已有数据库时可省略以复用保存的配置')
    imp.add_argument('--replace', action='store_true',
                     help='校验成功后完整替换目标数据库；使用前关闭其他数据库连接')
    query = commands.add_parser('query')
    query.add_argument('database'); query.add_argument('sql')
    query.add_argument('--params', default='{}'); query.add_argument('--limit', type=int, default=1000)
    query.add_argument('--timeout', type=float, default=5)
    args = parser.parse_args()
    try:
        if args.command == 'import':
            result = import_file(args.source, args.output,
                                 json.loads(Path(args.schema).read_text(encoding='utf-8')) if args.schema else None,
                                 replace=args.replace)
        else:
            result = query_file(args.database, args.sql, json.loads(args.params), args.limit, args.timeout)
        print(json.dumps(result, ensure_ascii=False, indent=2))
    except (ValueError, OSError, sqlite3.Error, KeyError) as exc:
        parser.exit(1, f'失败：{exc}\n')


if __name__ == '__main__':
    main()
